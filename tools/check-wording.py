# -*- coding: utf-8 -*-
"""
对外口径检查（CI 门禁）

目的：防止"自称复现/对标 Discord"一类表述、口径不一致与事实性错误再次混入对外交付物。
范围：面向外部的文本（README、docs/*.md、图源 *.puml、生成脚本）、代码注释，以及可选的 docx 成品。

用法：python tools/check-wording.py
退出码：0 = 通过；1 = 发现问题（CI 会因此失败）

规则维护：新增禁用表述请只改本文件下方的 FORBIDDEN 列表，不要散落到各处。
"""
import glob
import os
import re
import sys

# ---------- 仓库根定位 ----------
def find_root() -> str:
    p = os.path.dirname(os.path.abspath(__file__))
    while True:
        if os.path.exists(os.path.join(p, "README.md")) and os.path.isdir(os.path.join(p, "docs")):
            return p
        parent = os.path.dirname(p)
        if parent == p:
            raise RuntimeError("未找到仓库根目录")
        p = parent

ROOT = find_root()

TEXT_GLOBS = [
    "README.md",
    "docs/*.md",
    "docs/figures/*.puml",
    "docs/build/*.js",
    "tools/*.py",
]
CODE_GLOBS = [
    "src/main/java/**/*.java",
    "src/test/java/**/*.java",
    "client/src/**/*.ts",
    "client/src/**/*.tsx",
]
DOCX_GLOBS = ["docs/*.docx"]

# 有意保留旧文的备份文件（仅提示，不计失败）
EXEMPT_FILES = {"docs/毕业论文-原版备份.md"}
# 正当引用 Discord 的场景：行业举例 / 研究现状 / 参考文献 / 商标声明
ALLOW_LINE_PATTERNS = [
    r"参考文献",
    r"^\s*\[\d+\]\s*Discord",
    r"商标",
    r"以 Discord、",
    r"Discord 的网关",
    r"Discord 采用的",
    r"Discord 是这一领域的代表性产品",
]

FORBIDDEN = [
    (r"类\s?Discord", "自称体系为『类 Discord』"),
    (r"Discord\s?-?\s?like", "英文自称 Discord-like"),
    (r"Discord\s?Clone", "自称 Discord 克隆"),
    (r"复刻", "『复刻』——暗示照搬他人产品"),
    (r"标杆", "『标杆』——把对方立为照搬目标"),
    (r"对标", "『对标』"),
    (r"完全对齐", "『完全对齐』"),
    (r"与官方语义一致", "『与官方语义一致』"),
    (r"照搬|模仿|蓝本|克隆", "『照搬/模仿/蓝本/克隆』"),
]

MUST_HAVE = [
    ("docs/毕业论文.md", "基于 Spring Boot 与 React 的智能音视频实时通讯平台的设计与实现", "论文标题须为 A 题目"),
    ("docs/开题报告.md", "基于 Spring Boot 与 React 的智能音视频实时通讯平台的设计与实现", "开题报告题目须与论文一致"),
]

FACT_CHECKS = [
    ("docs/毕业论文.md", r"共\s*68\s*个自动化测试用例", "论文用例总数应为 68（45 后端 + 23 前端）", True),
    ("docs/毕业论文.md", r"(?<!6)63 个用例", "不应残留旧数字『63 个用例』", False),
]


def _walk_globs(globs):
    out = []
    for g in globs:
        out.extend(glob.glob(os.path.join(ROOT, g), recursive=True))
    return sorted(set(out))


def _allowed(line: str) -> bool:
    return any(re.search(p, line) for p in ALLOW_LINE_PATTERNS)


def scan_text_files(problems, notes):
    for path in _walk_globs(TEXT_GLOBS + CODE_GLOBS):
        rel = os.path.relpath(path, ROOT).replace("\\", "/")
        if rel == "tools/check-wording.py":
            continue  # 规则定义文件本身含禁用词，跳过
        try:
            text = open(path, encoding="utf-8", errors="ignore").read()
        except OSError:
            continue
        for lineno, line in enumerate(text.splitlines(), 1):
            if "wording-check:ignore" in line:
                continue  # 规则说明等确需引用禁用词的场合
            for pat, why in FORBIDDEN:
                if re.search(pat, line) and not _allowed(line):
                    item = (rel, lineno, why, line.strip()[:110])
                    (notes if rel in EXEMPT_FILES else problems).append(item)
    return problems, notes


def scan_docx(problems):
    """docx 成品同样纳入检查（需要 python-docx；缺失则跳过）"""
    try:
        import docx  # noqa
    except ImportError:
        print("  （未安装 python-docx，跳过 docx 扫描）")
        return
    for path in _walk_globs(DOCX_GLOBS):
        rel = os.path.relpath(path, ROOT).replace("\\", "/")
        if os.path.basename(path).startswith("~$"):
            continue
        try:
            d = docx.Document(path)
        except Exception:
            continue
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                for c in r.cells:
                    parts.append(c.text)
        for s in parts:
            if "wording-check:ignore" in s:
                continue
            for pat, why in FORBIDDEN:
                if re.search(pat, s) and not _allowed(s):
                    problems.append((rel, 0, why, s.strip()[:110]))


def main():
    problems, notes = [], []

    print("== 禁用表述（文本 + 代码注释）==")
    scan_text_files(problems, notes)
    scan_docx(problems)

    seen = set()
    uniq = []
    for it in problems:
        k = (it[0], it[1], it[2])
        if k not in seen:
            seen.add(k)
            uniq.append(it)
    problems = uniq

    if not problems:
        print("  ✅ 无")
    for rel, lineno, why, snippet in problems:
        loc = f"{rel}:{lineno}" if lineno else rel
        print(f"  ❌ {loc}  [{why}]\n     {snippet}")
    for rel, lineno, why, snippet in notes:
        print(f"  ℹ️  {rel}:{lineno}（备份文件，保留原样）")

    print("\n== 口径一致性 ==")
    for rel, needle, why in MUST_HAVE:
        p = os.path.join(ROOT, rel)
        ok = os.path.exists(p) and needle in open(p, encoding="utf-8", errors="ignore").read()
        print(f"  {'✅' if ok else '❌'} {rel}：{why}")
        if not ok:
            problems.append((rel, 0, why, needle))

    print("\n== 事实性核对 ==")
    for rel, pat, why, should_exist in FACT_CHECKS:
        p = os.path.join(ROOT, rel)
        if not os.path.exists(p):
            continue
        hit = re.search(pat, open(p, encoding="utf-8", errors="ignore").read())
        ok = bool(hit) == should_exist
        print(f"  {'✅' if ok else '❌'} {rel}：{why}")
        if not ok:
            problems.append((rel, 0, why, hit.group(0) if hit else "未匹配"))

    print(f"\n结论：{'❌ 存在问题 ' + str(len(problems)) + ' 处' if problems else '✅ 全部通过'}")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
